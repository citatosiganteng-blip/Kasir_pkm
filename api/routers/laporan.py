"""
KasirKu API — Laporan Router
Endpoint: dashboard summary, laporan harian
"""

from datetime import datetime, date, timedelta
from typing import Optional, List

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from io import BytesIO, StringIO
import csv

from database.db import db
from database.models import (
    Transaksi, TransaksiDetail, Barang, Pengeluaran, Pembelian, ReturPembelian,
)
from api.deps import get_current_user_payload
from api.schemas import DashboardSummaryOut
from utils.helpers import get_pembelian_status_display

router = APIRouter(prefix="/api/laporan", tags=["laporan"])


@router.get("/dashboard", response_model=DashboardSummaryOut)
def dashboard_summary(_: dict = Depends(get_current_user_payload)):
    """
    Ringkasan dashboard:
    - Total penjualan & transaksi hari ini
    - Total penjualan & transaksi bulan ini
    - Total pengeluaran hari ini
    - Laba bersih hari ini
    - Jumlah barang stok menipis
    - Top 5 produk terlaris hari ini
    """
    now = datetime.now()
    today_start = datetime(now.year, now.month, now.day, 0, 0, 0)
    today_end = datetime(now.year, now.month, now.day, 23, 59, 59)
    month_start = datetime(now.year, now.month, 1, 0, 0, 0)

    with db.get_session() as session:
        # — Hari ini —
        trx_hari_ini = session.query(Transaksi).filter(
            Transaksi.tanggal >= today_start,
            Transaksi.tanggal <= today_end,
            Transaksi.status == "selesai",
        ).all()

        total_hari_ini = sum(t.total for t in trx_hari_ini)
        jumlah_hari_ini = len(trx_hari_ini)

        # — Bulan ini —
        trx_bulan = session.query(Transaksi).filter(
            Transaksi.tanggal >= month_start,
            Transaksi.status == "selesai",
        ).all()
        total_bulan = sum(t.total for t in trx_bulan)
        jumlah_bulan = len(trx_bulan)

        # — Pengeluaran hari ini —
        pengeluaran_hari = session.query(Pengeluaran).filter(
            Pengeluaran.tanggal >= today_start,
            Pengeluaran.tanggal <= today_end,
        ).all()
        total_pengeluaran = sum(p.nominal for p in pengeluaran_hari)

        laba_bersih = total_hari_ini - total_pengeluaran

        # — Stok menipis —
        stok_menipis = session.query(Barang).filter(
            Barang.aktif == True,
            Barang.stok <= Barang.stok_min,
        ).count()

        # — Top 5 produk hari ini —
        top_produk_raw: dict[str, dict] = {}
        for t in trx_hari_ini:
            for d in t.detail:
                key = d.nama_barang
                if key not in top_produk_raw:
                    top_produk_raw[key] = {"nama": key, "qty": 0, "total": 0.0}
                top_produk_raw[key]["qty"] += d.qty
                top_produk_raw[key]["total"] += d.subtotal

        top_produk = sorted(
            top_produk_raw.values(),
            key=lambda x: x["qty"],
            reverse=True,
        )[:5]

        return DashboardSummaryOut(
            total_penjualan_hari_ini=total_hari_ini,
            jumlah_transaksi_hari_ini=jumlah_hari_ini,
            total_penjualan_bulan_ini=total_bulan,
            jumlah_transaksi_bulan_ini=jumlah_bulan,
            total_pengeluaran_hari_ini=total_pengeluaran,
            laba_bersih_hari_ini=laba_bersih,
            stok_menipis_count=stok_menipis,
            top_produk=top_produk,
        )


@router.get("/harian")
def laporan_harian(
    tanggal: Optional[date] = Query(None, description="Format: YYYY-MM-DD. Default: hari ini"),
    _: dict = Depends(get_current_user_payload),
):
    """Laporan detail per tanggal."""
    target = tanggal or date.today()
    start = datetime(target.year, target.month, target.day, 0, 0, 0)
    end = datetime(target.year, target.month, target.day, 23, 59, 59)

    with db.get_session() as session:
        transaksi_list = session.query(Transaksi).filter(
            Transaksi.tanggal >= start,
            Transaksi.tanggal <= end,
            Transaksi.status == "selesai",
        ).order_by(Transaksi.tanggal).all()

        pengeluaran_list = session.query(Pengeluaran).filter(
            Pengeluaran.tanggal >= start,
            Pengeluaran.tanggal <= end,
        ).all()

        total_penjualan = sum(t.total for t in transaksi_list)
        total_pengeluaran = sum(p.nominal for p in pengeluaran_list)

        by_metode: dict[str, float] = {}
        top_raw: dict[str, dict] = {}
        for t in transaksi_list:
            by_metode[t.metode_bayar] = by_metode.get(t.metode_bayar, 0) + t.total
            for d in t.detail:
                k=d.nama_barang
                row=top_raw.setdefault(k,{"nama":k,"qty":0,"total":0.0})
                row["qty"] += d.qty
                row["total"] += d.subtotal
        top_produk=sorted(top_raw.values(),key=lambda x:x["qty"],reverse=True)[:10]

        return {
            "tanggal": str(target),
            "jumlah_transaksi": len(transaksi_list),
            "total_penjualan": total_penjualan,
            "total_pengeluaran": total_pengeluaran,
            "laba_bersih": total_penjualan - total_pengeluaran,
            "penjualan_per_metode": by_metode,
            "top_produk": top_produk,
            "transaksi": [
                {
                    "id": t.id,
                    "no_invoice": t.no_invoice,
                    "waktu": t.tanggal.strftime("%H:%M"),
                    "kasir": t.kasir.username if t.kasir else "-",
                    "total": t.total,
                    "metode": t.metode_bayar,
                }
                for t in transaksi_list
            ],
            "pengeluaran": [
                {
                    "kategori": p.kategori,
                    "deskripsi": p.deskripsi,
                    "nominal": p.nominal,
                }
                for p in pengeluaran_list
            ],
        }


@router.get('/export.csv')
def export_csv(tanggal_mulai: Optional[date]=Query(None), tanggal_selesai: Optional[date]=Query(None), _:dict=Depends(get_current_user_payload)):
    start=tanggal_mulai or date.today(); end=tanggal_selesai or start
    with db.get_session() as s:
        from database.models import TransaksiDetail
        from sqlalchemy import func
        rows=s.query(TransaksiDetail.nama_barang, Barang.kategori, func.sum(TransaksiDetail.qty), func.avg(TransaksiDetail.harga), func.sum(TransaksiDetail.subtotal)).join(Transaksi).outerjoin(Barang, Barang.id==TransaksiDetail.barang_id).filter(Transaksi.status=='selesai', Transaksi.tanggal>=datetime.combine(start,datetime.min.time()), Transaksi.tanggal<=datetime.combine(end,datetime.max.time())).group_by(TransaksiDetail.nama_barang, Barang.kategori).order_by(func.sum(TransaksiDetail.qty).desc()).all()
    out=StringIO(); w=csv.writer(out); w.writerow(['No','Nama Produk','Kategori','Terjual','Harga Rata-rata','Total Pendapatan'])
    for i,r in enumerate(rows,1): w.writerow([i,r[0],r[1] or '-',int(r[2] or 0),float(r[3] or 0),float(r[4] or 0)])
    return StreamingResponse(iter([out.getvalue().encode('utf-8-sig')]),media_type='text/csv',headers={'Content-Disposition':f'attachment; filename=laporan_penjualan_{start}_{end}.csv'})

@router.get('/export.xlsx')
def export_xlsx(tanggal_mulai: Optional[date]=Query(None), tanggal_selesai: Optional[date]=Query(None), _:dict=Depends(get_current_user_payload)):
    start=tanggal_mulai or date.today(); end=tanggal_selesai or start
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except Exception as e:
        raise HTTPException(500,'Library openpyxl tidak tersedia')
    with db.get_session() as s:
        from database.models import TransaksiDetail
        from sqlalchemy import func
        from sqlalchemy.orm import joinedload, selectinload
        rows=s.query(TransaksiDetail.nama_barang, Barang.kategori, func.sum(TransaksiDetail.qty), func.avg(TransaksiDetail.harga), func.sum(TransaksiDetail.subtotal)).join(Transaksi).outerjoin(Barang, Barang.id==TransaksiDetail.barang_id).filter(Transaksi.status=='selesai', Transaksi.tanggal>=datetime.combine(start,datetime.min.time()), Transaksi.tanggal<=datetime.combine(end,datetime.max.time())).group_by(TransaksiDetail.nama_barang, Barang.kategori).order_by(func.sum(TransaksiDetail.qty).desc()).all()
        purchases=s.query(Pembelian).options(
            joinedload(Pembelian.supplier),
            selectinload(Pembelian.detail),
            selectinload(Pembelian.retur).selectinload(ReturPembelian.detail),
        ).filter(
            Pembelian.tanggal >= datetime.combine(start, datetime.min.time()),
            Pembelian.tanggal <= datetime.combine(end, datetime.max.time()),
        ).order_by(Pembelian.tanggal.asc()).all()
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='Laporan Penjualan'
    ws.append(['No','Nama Produk','Kategori','Terjual','Harga Rata-rata','Total Pendapatan'])
    for i,r in enumerate(rows,1): ws.append([i,r[0],r[1] or '-',int(r[2] or 0),float(r[3] or 0),float(r[4] or 0)])
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='2563EB'); c.alignment=Alignment(horizontal='center')
    po_ws=wb.create_sheet('Pembelian & Restock')
    store_name = db.get_setting("store_name", "KasirKu")
    po_ws.cell(1, 1, store_name).font = Font(bold=True, size=16, color="1E293B")
    po_ws.cell(2, 1, f"Periode: {start.strftime('%d/%m/%Y')} s/d {end.strftime('%d/%m/%Y')}").font = Font(size=11, italic=True, color="64748B")
    po_headers = [
        "No PO", "No Faktur Vendor", "Tanggal", "Supplier", "Nama Barang",
        "Qty Restock", "Harga Beli", "Subtotal Item", "Total Faktur",
        "Sudah Dibayar", "Sisa Hutang", "Status", "Jatuh Tempo",
        "Keterangan Retur",
    ]
    po_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    for col_idx, heading in enumerate(po_headers, 1):
        cell = po_ws.cell(4, col_idx, heading)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = po_border
    for p in purchases:
        detail=p.detail or [None]
        status_text, return_note = get_pembelian_status_display(p)
        for item in detail:
            po_ws.append([
                p.no_po, p.no_faktur, p.tanggal, p.supplier.nama if p.supplier else "Umum",
                item.nama_barang if item else "-", item.qty if item else 0,
                item.harga_beli if item else 0, item.subtotal if item else 0,
                p.total or 0, p.sudah_dibayar or 0,
                max(0, (p.total or 0) - (p.sudah_dibayar or 0)),
                status_text, p.jatuh_tempo, return_note or "-",
            ])
            row_num = po_ws.max_row
            for c in po_ws[row_num]:
                c.border = po_border
                if c.column in (3, 13) and c.value:
                    c.number_format = "dd/mm/yyyy"
                elif c.column == 6:
                    c.number_format = "#,##0.##"
                elif c.column in (7, 8, 9, 10, 11):
                    c.number_format = '"Rp" #,##0.00'
                if c.column in (6, 7, 8, 9, 10, 11):
                    c.alignment = Alignment(horizontal="right", vertical="top")
                elif c.column == 14:
                    c.alignment = Alignment(vertical="top", wrap_text=True)
                else:
                    c.alignment = Alignment(vertical="top")

            status_cell = po_ws.cell(row_num, 12)
            if p.retur:
                status_cell.font = Font(color="B45309", bold=True)
                status_cell.fill = PatternFill("solid", fgColor="FEF3C7")
                po_ws.cell(row_num, 14).font = Font(color="B45309")
                po_ws.cell(row_num, 14).fill = PatternFill("solid", fgColor="FEF3C7")
            elif p.status_bayar in ("tempo", "cicil"):
                status_cell.font = Font(color="DC2626", bold=True)
                status_cell.fill = PatternFill("solid", fgColor="FEE2E2")
            elif p.status_bayar == "lunas":
                status_cell.font = Font(color="059669", bold=True)

    for col,width in enumerate([18,22,18,24,28,14,18,20,20,20,20,34,18,55],1):
        po_ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width=width
    po_ws.freeze_panes='A5'; po_ws.auto_filter.ref=f'A4:N{max(4,po_ws.max_row)}'
    po_ws.row_dimensions[4].height = 32
    bio=BytesIO(); wb.save(bio); bio.seek(0)
    return StreamingResponse(bio,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename=laporan_penjualan_{start}_{end}.xlsx'})

@router.get('/export-tax.xlsx')
def export_tax_xlsx(tanggal_mulai: Optional[date]=Query(None), tanggal_selesai: Optional[date]=Query(None), _:dict=Depends(get_current_user_payload)):
    start=tanggal_mulai or date.today(); end=tanggal_selesai or start
    try:
        import openpyxl
        from sqlalchemy import func
    except Exception:
        raise HTTPException(500,'Library export tidak tersedia')
    with db.get_session() as s:
        trx=s.query(Transaksi).filter(Transaksi.status=='selesai',Transaksi.tanggal>=datetime.combine(start,datetime.min.time()),Transaksi.tanggal<=datetime.combine(end,datetime.max.time())).all()
        pemb=s.query(Pembelian).filter(Pembelian.tanggal>=datetime.combine(start,datetime.min.time()),Pembelian.tanggal<=datetime.combine(end,datetime.max.time())).all()
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='Rekap Pajak'
    ws.append(['Jenis','Nomor','Tanggal','DPP','PPN %','PPN Nominal','PPh %','PPh Nominal','Total'])
    for t in trx: ws.append(['Penjualan',t.no_invoice,t.tanggal.strftime('%Y-%m-%d'),t.dpp,t.ppn_persen,t.ppn_nominal,t.pph_persen,t.pph_nominal,t.total])
    for p in pemb: ws.append(['Pembelian',p.no_po,p.tanggal.strftime('%Y-%m-%d'),p.dpp,p.ppn_persen,p.ppn_nominal,p.pph_persen,p.pph_nominal,p.total])
    for c in ws[1]: c.font=openpyxl.styles.Font(bold=True,color='FFFFFF'); c.fill=openpyxl.styles.PatternFill('solid',fgColor='2563EB')
    bio=BytesIO(); wb.save(bio); bio.seek(0)
    return StreamingResponse(bio,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename=rekap_pajak_ppn_{start}_{end}.xlsx'})
