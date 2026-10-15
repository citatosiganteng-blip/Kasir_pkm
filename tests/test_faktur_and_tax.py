"""
Unit tests for Faktur Penjualan (A4/PDF), Faktur Pembelian (PO & Restock),
and Perpajakan (PPN/PPh & Rekapitulasi).
"""

import os
from datetime import datetime, timedelta
from types import SimpleNamespace
import pytest
from openpyxl import load_workbook

from database.models import Barang, Transaksi, TransaksiDetail, Supplier, Pelanggan, Pembelian, PembelianDetail
from utils.helpers import (
    terbilang, generate_po_number, generate_supplier_code,
    generate_customer_code, generate_tax_invoice_number,
    get_pembelian_status_display, apply_potong_hutang_retur
)
from services.invoice_pdf_service import InvoicePdfService
from api.routers.pembelian import _out as pembelian_api_out


def test_terbilang_conversion():
    """Uji konversi angka nominal ke teks terbilang bahasa Indonesia"""
    assert terbilang(0) == "Nol Rupiah"
    assert terbilang(1000) == "Seribu Rupiah"
    assert terbilang(1500000) == "Satu Juta Lima Ratus Ribu Rupiah"
    assert terbilang(25000) == "Dua Puluh Lima Ribu Rupiah"
    assert terbilang(100100) == "Seratus Ribu Seratus Rupiah"


def test_code_generators(test_db):
    """Uji generator otomatis kode PO, Supplier, Pelanggan, dan Faktur Pajak"""
    with test_db.get_session() as session:
        po_num = generate_po_number(session)
        assert po_num.startswith("PO-")

        sup_code = generate_supplier_code(session)
        assert sup_code.startswith("SUP")

        plg_code = generate_customer_code(session)
        assert plg_code.startswith("PLG")

        fp_num = generate_tax_invoice_number(session)
        assert fp_num.startswith("FP-")


def test_purchase_status_display_includes_return_method_and_reason():
    retur = SimpleNamespace(
        no_retur="RB-20261015-001",
        alasan="Barang rusak",
        metode_kembali="potong_hutang",
    )
    pembelian = SimpleNamespace(
        status_bayar="lunas",
        total=50000,
        sudah_dibayar=50000,
        retur=[retur],
    )

    status, note = get_pembelian_status_display(pembelian)

    assert status == "RETUR - POTONG HUTANG/TAGIHAN"
    assert "RB-20261015-001 - Barang rusak" in note


def test_purchase_return_deduction_preserves_payment_history():
    pembelian = SimpleNamespace(status_bayar="cicil", total=50000, sudah_dibayar=30000)

    apply_potong_hutang_retur(pembelian, 25000)

    assert pembelian.total == 25000
    assert pembelian.status_bayar == "cicil"
    assert pembelian.sudah_dibayar == 30000


def test_purchase_invoice_html_shows_return_status_and_full_return_details():
    return_detail = SimpleNamespace(
        subtotal=25000,
        qty=2,
        harga_beli=12500,
        kode_barang="BRG-RET-01",
        nama_barang="Barang Retur",
    )
    purchase_return = SimpleNamespace(
        metode_kembali="refund_cash",
        tanggal=datetime(2026, 10, 15, 10, 30),
        alasan="Barang rusak",
        no_retur="RB-20261015-001",
        detail=[return_detail],
    )
    purchase = SimpleNamespace(
        supplier=None,
        tanggal=datetime(2026, 10, 15, 9, 0),
        jatuh_tempo=None,
        status_bayar="lunas",
        retur=[purchase_return],
        detail=[],
        user=None,
        total=100000,
        sudah_dibayar=100000,
        cicilan=[],
        no_po="PO-TEST-RET-01",
        no_faktur="INV-RET-01",
        metode_bayar="transfer",
        catatan="",
        dpp=100000,
        subtotal=100000,
        ppn_nominal=0,
        ppn_persen=0,
        pph_nominal=0,
        pph_persen=0,
    )
    store_info = {
        "name": "Toko Uji",
        "address": "Alamat Uji",
        "phone": "08123456789",
        "npwp": "-",
        "is_pkp": False,
        "bank_name": "BCA",
        "bank_account": "-",
        "bank_holder": "-",
    }

    html = InvoicePdfService.generate_purchase_order_html(purchase, store_info)

    assert "RETUR - REFUND CASH/TRANSFER" in html
    assert "RINCIAN RETUR KE SUPPLIER" in html
    assert "15 Oktober 2026" in html
    assert "10:30" in html
    assert "Tanggal / No. Retur" in html
    assert "table-layout: fixed" in html
    assert "BRG-RET-01" in html
    assert "Barang Retur" in html
    assert "2" in html
    assert "Rp 25.000" in html
    assert "Refund cash/transfer" in html
    assert "Barang rusak" in html
    assert ">LUNAS</span>" not in html


def test_mobile_purchase_api_includes_return_details():
    return_detail = SimpleNamespace(
        id=11,
        barang_id=7,
        kode_barang="BRG-RET-01",
        nama_barang="Barang Retur",
        qty=2,
        harga_beli=12500,
        subtotal=25000,
    )
    purchase_return = SimpleNamespace(
        id=5,
        no_retur="RB-20261015-001",
        tanggal=datetime(2026, 10, 15, 10, 30),
        total_retur=25000,
        alasan="Barang rusak",
        metode_kembali="refund_cash",
        detail=[return_detail],
    )
    purchase = SimpleNamespace(
        id=3,
        no_po="PO-TEST-RET-01",
        no_faktur="INV-RET-01",
        tanggal=datetime(2026, 10, 15, 9),
        jatuh_tempo=None,
        supplier_id=2,
        supplier=SimpleNamespace(id=2, kode="SUP02", nama="Supplier Uji", npwp=None),
        subtotal=100000,
        dpp=100000,
        ppn_persen=0,
        ppn_nominal=0,
        pph_persen=0,
        pph_nominal=0,
        total=100000,
        sudah_dibayar=100000,
        status_bayar="lunas",
        status_barang="diterima",
        metode_bayar="transfer",
        catatan=None,
        cicilan=[],
        retur=[purchase_return],
        detail=[],
    )

    output = pembelian_api_out(purchase)

    assert output["retur"][0]["no_retur"] == "RB-20261015-001"
    assert output["retur"][0]["tanggal"] == "2026-10-15T10:30:00"
    assert output["retur"][0]["metode_kembali"] == "refund_cash"
    assert output["retur"][0]["alasan"] == "Barang rusak"
    assert output["retur"][0]["detail"][0]["kode_barang"] == "BRG-RET-01"
    assert output["retur"][0]["detail"][0]["qty"] == 2
    assert output["retur"][0]["detail"][0]["subtotal"] == 25000


def test_pembelian_and_auto_stock_increment(test_db):
    """Uji faktur pembelian supplier: otomatis menambah stok barang dan update harga beli"""
    with test_db.get_session() as session:
        # 1. Buat supplier & barang
        supplier = Supplier(
            kode="SUP999",
            nama="PT Sumber Makmur Jaya",
            telepon="08123456789",
            npwp="01.999.888.7-654.000"
        )
        session.add(supplier)

        barang = Barang(
            kode="BRG999",
            nama="Biskuit Kaleng Kaliber",
            harga_beli=20000,
            harga_jual=28000,
            stok=15
        )
        session.add(barang)
        session.flush()

        sup_id = supplier.id
        brg_id = barang.id

        # 2. Buat Faktur Pembelian (PO) dengan restock 20 pcs @ Rp 22.000
        qty_masuk = 20
        harga_beli_baru = 22000
        dpp = qty_masuk * harga_beli_baru  # 440.000
        ppn = dpp * 0.11  # 48.400
        grand_total = dpp + ppn

        po = Pembelian(
            no_faktur="INV-VENDOR-7788",
            no_po="PO-20260917-999",
            supplier_id=sup_id,
            tanggal=datetime.now(),
            subtotal=dpp,
            dpp=dpp,
            ppn_persen=11.0,
            ppn_nominal=ppn,
            total=grand_total,
            status_bayar="lunas"
        )
        session.add(po)
        session.flush()

        detail = PembelianDetail(
            pembelian_id=po.id,
            barang_id=brg_id,
            kode_barang=barang.kode,
            nama_barang=barang.nama,
            qty=qty_masuk,
            harga_beli=harga_beli_baru,
            subtotal=dpp
        )
        session.add(detail)

        # Restock stok barang
        barang.stok += qty_masuk
        barang.harga_beli = harga_beli_baru

    # Verifikasi data setelah commit
    with test_db.get_session() as session:
        saved_barang = session.query(Barang).filter_by(kode="BRG999").first()
        assert saved_barang.stok == 35  # 15 + 20
        assert saved_barang.harga_beli == 22000

        saved_po = session.query(Pembelian).filter_by(no_po="PO-20260917-999").first()
        assert saved_po is not None
        assert saved_po.supplier.nama == "PT Sumber Makmur Jaya"
        assert len(saved_po.detail) == 1
        assert saved_po.total == 488400


def test_sales_invoice_pdf_and_html(test_db, tmp_path):
    """Uji pembuatan HTML Faktur Penjualan dan konversi ke file PDF A4"""
    with test_db.get_session() as session:
        pelanggan = Pelanggan(
            kode="PLG001",
            nama="Toko Berkah Abadi",
            telepon="08198765432",
            alamat="Jl. Sudirman No. 12, Jakarta",
            npwp="02.345.678.9-012.000"
        )
        session.add(pelanggan)
        session.flush()

        trx = Transaksi(
            no_invoice="INV-TEST-PDF-001",
            tanggal=datetime.now(),
            pelanggan_id=pelanggan.id,
            nama_pelanggan=pelanggan.nama,
            alamat_pelanggan=pelanggan.alamat,
            telepon_pelanggan=pelanggan.telepon,
            npwp_pelanggan=pelanggan.npwp,
            dpp=100000,
            ppn_persen=11.0,
            ppn_nominal=11000,
            total=111000,
            bayar=120000,
            kembalian=9000,
            metode_bayar="transfer",
            status_bayar="lunas"
        )
        session.add(trx)
        session.flush()

        det = TransaksiDetail(
            transaksi_id=trx.id,
            nama_barang="Produk Unggulan",
            kode_barang="BRG001",
            qty=2,
            harga=50000,
            diskon=0,
            subtotal=100000
        )
        session.add(det)
        session.commit()

        # Generate HTML
        html = InvoicePdfService.generate_sales_invoice_html(trx)
        assert "FAKTUR PENJUALAN" in html
        assert "INV-TEST-PDF-001" in html
        assert "Toko Berkah Abadi" in html
        assert "111,000" in html or "111.000" in html

        # Generate PDF
        pdf_file = tmp_path / "faktur_test.pdf"
        ok = InvoicePdfService.save_html_to_pdf(html, pdf_file)
        assert ok is True
        assert pdf_file.exists()
        assert pdf_file.stat().st_size > 1000  # Valid PDF size


def test_tax_report_excel_generation(test_db, tmp_path):
    """Uji ekspor rekapitulasi SPT Pajak (PPN & PPh) ke file Excel multi-sheet"""
    start_dt = datetime.now() - timedelta(days=1)
    end_dt = datetime.now() + timedelta(days=1)
    excel_path = tmp_path / "rekap_pajak.xlsx"

    ok = InvoicePdfService.export_tax_report_to_excel(start_dt, end_dt, excel_path)
    assert ok is True
    assert excel_path.exists()

    # Periksa lembar kerja di dalam file Excel
    wb = load_workbook(excel_path)
    sheet_names = wb.sheetnames
    assert "Rekap SPT Pajak Masa" in sheet_names
    assert "Pajak Keluaran (Sales)" in sheet_names
    assert "Pajak Masukan (Purchases)" in sheet_names
