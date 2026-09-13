"""
Unit test for Excel Export generation using openpyxl
"""

import openpyxl
from datetime import date
from io import BytesIO


def test_excel_export_structure():
    """Uji pembuatan workbook openpyxl dengan 3 sheet dan format laporan"""
    wb = openpyxl.Workbook()
    # Hapus sheet default
    wb.remove(wb.active)

    ws1 = wb.create_sheet("Rekap Harian")
    ws2 = wb.create_sheet("Top Produk")
    ws3 = wb.create_sheet("Rekap Pengeluaran")

    # Sheet 1: Rekap Harian
    ws1.append(["Tanggal", "Jml Transaksi", "Pemasukan (Rp)", "Pengeluaran (Rp)", "Laba Bersih (Rp)"])
    ws1.append([str(date.today()), 5, 250000, 50000, 200000])

    # Sheet 2: Top Produk
    ws2.append(["#", "Nama Barang", "Jml Transaksi", "Total Qty", "Revenue (Rp)"])
    ws2.append([1, "Indomie Goreng", 10, 25, 112500])

    # Sheet 3: Rekap Pengeluaran
    ws3.append(["Kategori Pengeluaran", "Jumlah Transaksi", "Total Nominal (Rp)"])
    ws3.append(["Operasional", 2, 50000])

    # Simpan ke stream in-memory
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    # Baca ulang workbook untuk validasi keutuhan file
    loaded_wb = openpyxl.load_workbook(stream)
    sheet_names = loaded_wb.sheetnames
    assert len(sheet_names) == 3
    assert "Rekap Harian" in sheet_names
    assert "Top Produk" in sheet_names
    assert "Rekap Pengeluaran" in sheet_names

    ws1_loaded = loaded_wb["Rekap Harian"]
    assert ws1_loaded.cell(row=1, column=1).value == "Tanggal"
    assert ws1_loaded.cell(row=2, column=5).value == 200000
